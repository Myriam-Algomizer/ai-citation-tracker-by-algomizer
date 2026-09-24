#!/usr/bin/env python3
"""Build the citation workbook and summary.json from parsed runs and enriched pages.

Usage:
  python build_sheet.py --work workdir/ --config config.json --out report.xlsx

Reads workdir/runs.json and workdir/pages.json (written by parse_logs.py, with
source / type / mentions filled in by Claude). Writes the .xlsx and
workdir/summary.json, and prints the figures for the chat summary.
"""
import argparse, itertools, json, math, os, random, re
from collections import Counter, defaultdict
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

MIN_RUNS_BRANDS, MIN_RUNS_SOURCES = 7, 8
TITLE_F = Font(name="Arial", size=12, bold=True, color="FFFFFF"); TITLE_FILL = PatternFill("solid", fgColor="1A1A2E")
HEAD_F = Font(name="Arial", size=11, bold=True, color="FFFFFF"); HEAD_FILL = PatternFill("solid", fgColor="1B2A4A")
F = Font(name="Arial", size=10); NOTE = Font(name="Arial", size=9, italic=True)
LINK = Font(name="Arial", size=10, color="1155CC", underline="single")
BRAND_F = Font(name="Arial", size=10, bold=True, color="5566F4")
YELLOW = PatternFill("solid", fgColor="FFF2CC"); ORANGE = PatternFill("solid", fgColor="FCE4D6"); BLUE = PatternFill("solid", fgColor="E3E6FD")


def wilson(k, n, z=1.96):
    if n == 0: return (0.0, 0.0)
    p = k / n; den = 1 + z * z / n; mid = p + z * z / (2 * n)
    half = z * math.sqrt(p * (1 - p) / n + z * z / (4 * n * n))
    return (max(0.0, (mid - half) / den), min(1.0, (mid + half) / den))


def jaccard_mean(sets):
    vals = []
    for a, b in itertools.combinations(sets, 2):
        if a or b: vals.append(len(a & b) / len(a | b))
    return round(sum(vals) / len(vals), 2) if vals else None


def boot_share(runs, dom_of, target, B=1000, seed=7):
    rnd = random.Random(seed); out = []
    for _ in range(B):
        s = [rnd.choice(runs) for _ in runs]
        tot = sum(c["count"] for r in s for c in r["citations"])
        hit = sum(c["count"] for r in s for c in r["citations"] if dom_of(c["url"]) == target)
        out.append(hit / tot if tot else 0.0)
    out.sort(); return (out[int(0.025 * B)], out[int(0.975 * B) - 1])


def owner(dom, cfg):
    def match(d, doms): return any(d == x or d.endswith("." + x) for x in doms)
    if match(dom, cfg["brand"].get("domains", [])): return "brand"
    for c in cfg.get("competitors", []):
        if match(dom, c.get("domains", [])): return c["name"]
    return None


def safe_sheet(name, used):
    n = re.sub(r"[\[\]\*\?/\\:]", " ", name)[:31].strip() or "LLM"
    while n in used: n = n[:28] + "_" + str(len(used))
    used.add(n); return n


def scope_stats(runs, cfg, dom_of):
    n = len(runs); ents = [cfg["brand"]["name"]] + [c["name"] for c in cfg.get("competitors", [])]
    brands = {}
    for e in ents:
        k = sum(1 for r in runs if e in r["brands"]); lo, hi = wilson(k, n)
        pos = [r["brands"].index(e) + 1 for r in runs if e in r["brands"]]
        brands[e] = {"answers": k, "runs": n, "rate": round(k / n, 3) if n else 0, "low": round(lo, 3), "high": round(hi, 3),
                     "named_first": sum(1 for p in pos if p == 1), "avg_position": round(sum(pos) / len(pos), 1) if pos else None}
    dc, dr, pr = Counter(), defaultdict(set), defaultdict(set)
    for r in runs:
        for c in r["citations"]:
            d = dom_of(c["url"]); dc[d] += c["count"]; dr[d].add(r["id"]); pr[c["url"]].add(r["id"])
    tot = sum(dc.values()); domains = []
    for d, k in dc.most_common(10):
        row = {"domain": d, "citations": k, "share": round(k / tot, 3), "runs": len(dr[d]), "appearance": round(len(dr[d]) / n, 3)}
        if n >= MIN_RUNS_SOURCES:
            lo, hi = boot_share(runs, dom_of, d); row["share_low"], row["share_high"] = round(lo, 3), round(hi, 3)
        domains.append(row)
    cited = [set(c["url"] for c in r["citations"]) for r in runs]
    return {"runs": n, "total_citations": tot, "unique_pages": len(pr), "brands": brands, "top_domains": domains,
            "pages_seen_in_one_run_only": sum(1 for v in pr.values() if len(v) == 1),
            "runs_without_sources": sum(1 for s in cited if not s),
            "source_stability_jaccard": jaccard_mean([s for s in cited if s]),
            "brand_stability_jaccard": jaccard_mean([set(r["brands"]) for r in runs])}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--work", required=True); ap.add_argument("--config", required=True); ap.add_argument("--out", required=True)
    a = ap.parse_args()
    cfg = json.load(open(a.config))
    data = json.load(open(os.path.join(a.work, "runs.json"))); runs, prompts = data["runs"], data["prompts"]
    pages = {p["url"]: p for p in json.load(open(os.path.join(a.work, "pages.json")))}
    brand = cfg["brand"]["name"]; prepared = cfg.get("prepared_by", "")
    dom_of = lambda u: pages[u]["domain"]
    multi_prompt = len(prompts) > 1

    stamps = sorted(datetime.fromisoformat(r["timestamp"]) for r in runs if r["timestamp"])
    if stamps:
        a0, a1 = stamps[0].strftime("%b %d, %Y"), stamps[-1].strftime("%b %d, %Y")
        dates = a0 if a0 == a1 else f"{a0} to {a1}"
    else:
        dates = "dates unknown"

    llms = [l for l, _ in Counter(r["llm"] for r in runs).most_common()]
    cell = Counter((r["prompt_id"], r["llm"]) for r in runs)
    warnings = []
    for (p, l), k in sorted(cell.items()):
        if k < MIN_RUNS_SOURCES:
            warnings.append(f"{p} on {l}: {k} run(s). Minimum is {MIN_RUNS_BRANDS} for brand figures and {MIN_RUNS_SOURCES} for source figures.")

    wb = Workbook(); wb.remove(wb.active); used = set()
    answers_name = "Answers"; used.add(answers_name)
    row_of = {r["id"]: 4 + i for i, r in enumerate(runs)}

    def cit_tab(label, sel, is_all):
        ws = wb.create_sheet(safe_sheet(label, used)); n = len(sel)
        cit, rc, kinds, mods, prs = Counter(), defaultdict(list), {}, defaultdict(list), defaultdict(list)
        for r in sel:
            for c in r["citations"]:
                u = c["url"]; cit[u] += c["count"]; rc[u].append(r["id"]); kinds.setdefault(u, c["kind"])
                if r["llm"] not in mods[u]: mods[u].append(r["llm"])
                if r["prompt_id"] not in prs[u]: prs[u].append(r["prompt_id"])
        hdr = ["Rank", "Page URL", "Domain", "Title", "Source", "Type", "Citations", "Citation %", "Appearance rate", "Cited in runs (click for answer)", "Kind"]
        wid = [7, 58, 22, 46, 13, 16, 10, 11, 13, 20, 19]
        if is_all: hdr.append("Models"); wid.append(22)
        if multi_prompt: hdr.append("Prompts"); wid.append(12)
        hdr += [f"{brand} mentioned", "Competitors mentioned"]; wid += [14, 36]
        ncol = len(hdr); last = ws.cell(1, ncol).column_letter
        ws.merge_cells(f"A1:{last}1")
        bits = [brand, "AI Citation Analysis", label, dates, f"{n} runs", f"{len(set(r['prompt_id'] for r in sel))} prompt(s)"]
        if prepared: bits.append(f"Prepared by {prepared}")
        ws["A1"] = " | ".join(bits); ws["A1"].font = TITLE_F; ws["A1"].fill = TITLE_FILL
        ws["A1"].alignment = Alignment(vertical="center"); ws.row_dimensions[1].height = 24
        legend = f"Yellow = page names a competitor. Orange = competitor's own site. Blue = {brand}'s own site or a page naming {brand}. Appearance rate = share of runs that cited the page."
        tab_warn = [w for w in warnings if is_all or f" on {label}:" in w]
        if tab_warn: legend += "   WARNING: " + " ".join(tab_warn)
        ws["A2"] = legend; ws["A2"].font = NOTE
        for i, h in enumerate(hdr, 1):
            c = ws.cell(3, i, h); c.font = HEAD_F; c.fill = HEAD_FILL; c.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[c.column_letter].width = wid[i - 1]
        ws.row_dimensions[3].height = 32; ws.freeze_panes = "C4"
        total = sum(cit.values())
        order = sorted(cit, key=lambda u: (-cit[u], -len(rc[u]), u))
        for k, u in enumerate(order, 1):
            p = pages[u]; row = 3 + k
            vals = [f"#{k}", u, p["domain"], p.get("title", ""), p.get("source", ""), p.get("type", ""), cit[u],
                    cit[u] / total if total else 0, len(rc[u]) / n if n else 0, ", ".join(rc[u]), kinds[u]]
            if is_all: vals.append(", ".join(mods[u]))
            if multi_prompt: vals.append(", ".join(prs[u]))
            bm, cm = p.get("brand_mentioned", "not checked"), p.get("competitors_mentioned", "not checked")
            vals += [bm, cm]
            for j, v in enumerate(vals, 1): ws.cell(row, j, v).font = F
            ws.cell(row, 2).hyperlink = u; ws.cell(row, 2).font = LINK
            ws.cell(row, 8).number_format = "0.0%"; ws.cell(row, 9).number_format = "0%"
            lc = ws.cell(row, 10); lc.hyperlink = f"#'{answers_name}'!A{row_of[rc[u][0]]}"; lc.font = LINK
            own = owner(p["domain"], cfg)
            if own == "brand": ws.cell(row, 3).fill = BLUE
            elif own: ws.cell(row, 3).fill = ORANGE
            if cm and not cm.lower().startswith(("not checked", "none", "no", "unknown")): ws.cell(row, ncol).fill = YELLOW
            if str(bm).lower().startswith("yes"): ws.cell(row, ncol - 1).font = BRAND_F
        if order: ws.auto_filter.ref = f"A3:{last}{3 + len(order)}"

    cit_tab("All LLMs", runs, True)
    for l in llms: cit_tab(l, [r for r in runs if r["llm"] == l], False)

    if multi_prompt:
        ws = wb.create_sheet(safe_sheet("Prompts", used))
        hdr = ["Prompt", "Text"] + llms
        for i, h in enumerate(hdr, 1):
            c = ws.cell(1, i, h); c.font = HEAD_F; c.fill = HEAD_FILL
        ws.column_dimensions["A"].width = 9; ws.column_dimensions["B"].width = 90
        for i, p in enumerate(prompts, 2):
            ws.cell(i, 1, p["id"]).font = F; ws.cell(i, 2, p["text"]).font = F
            for j, l in enumerate(llms, 3): ws.cell(i, j, cell.get((p["id"], l), 0)).font = F

    ws = wb.create_sheet(answers_name)
    hdr = ["Run", "Prompt", "Timestamp", "LLM", "Model", "Mode", "Brands named, in order", "New names outside your list", "Sources", "Source file", "Full answer"]
    wid = [7, 9, 24, 11, 18, 18, 34, 22, 9, 38, 120]
    ws.merge_cells("A1:K1"); ws["A1"] = "Full answers, one row per run. The citation tabs link here."
    ws["A1"].font = TITLE_F; ws["A1"].fill = TITLE_FILL; ws.row_dimensions[1].height = 24
    ws["A2"] = "   ".join(f"{p['id']}: {p['text']}" for p in prompts)[:2000]; ws["A2"].font = NOTE
    for i, h in enumerate(hdr, 1):
        c = ws.cell(3, i, h); c.font = HEAD_F; c.fill = HEAD_FILL; ws.column_dimensions[c.column_letter].width = wid[i - 1]
    ws.freeze_panes = "B4"
    for r in runs:
        row = row_of[r["id"]]; ans = r["answer"]
        if len(ans) > 32000: ans = ans[:32000] + "\n[cut at the Excel cell limit. Full text is in the log file.]"
        vals = [r["id"], r["prompt_id"], r["timestamp_raw"], r["llm"], r["model"], r["mode"], "; ".join(r["brands"]),
                "; ".join(r["new_names"]), sum(c["count"] for c in r["citations"]), r["file"], ans]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row, j, v); c.font = F; c.alignment = Alignment(wrap_text=(j in (7, 8, 11)), vertical="top")
        ws.row_dimensions[row].height = 180
    wb.save(a.out)

    # ---------- summary ----------
    summary = {"brand": brand, "dates": dates, "prompts": prompts, "runs_per_prompt_and_llm": {f"{p} x {l}": k for (p, l), k in sorted(cell.items())},
               "warnings": warnings, "all": scope_stats(runs, cfg, dom_of), "by_llm": {l: scope_stats([r for r in runs if r["llm"] == l], cfg, dom_of) for l in llms}}
    sets = {l: set(c["url"] for r in runs if r["llm"] == l for c in r["citations"]) for l in llms}
    summary["shared_pages_between_llms"] = {f"{x} & {y}": len(sets[x] & sets[y]) for x, y in itertools.combinations(llms, 2)}
    new = Counter(nm for r in runs for nm in r["new_names"])
    summary["new_names_outside_list"] = dict(new.most_common(15))
    # brand vs competitor separation, only where the run count supports it
    sep = []
    for l, st in summary["by_llm"].items():
        if st["runs"] < MIN_RUNS_BRANDS: continue
        b = st["brands"][brand]
        for c, v in st["brands"].items():
            if c != brand and v["low"] <= b["high"] and v["high"] >= b["low"] and v["rate"] != b["rate"]:
                sep.append(f"{l}: {brand} ({b['rate']:.0%}) and {c} ({v['rate']:.0%}) have overlapping ranges. Treat them as level.")
    summary["overlap_notes"] = sep
    json.dump(summary, open(os.path.join(a.work, "summary.json"), "w"), indent=1, ensure_ascii=False)

    print(f"saved {a.out}")
    for l, st in [("ALL", summary["all"])] + list(summary["by_llm"].items()):
        print(f"\n[{l}] {st['runs']} runs, {st['total_citations']} citations, {st['unique_pages']} pages, "
              f"{st['pages_seen_in_one_run_only']} pages seen in one run only, source stability {st['source_stability_jaccard']}, brand stability {st['brand_stability_jaccard']}")
        for e, v in st["brands"].items():
            print(f"   {e}: {v['answers']} of {v['runs']} answers ({v['rate']:.0%}, range {v['low']:.0%} to {v['high']:.0%}), named first {v['named_first']}x, avg position {v['avg_position']}")
        for d in st["top_domains"][:6]:
            rng = f", range {d['share_low']:.0%} to {d['share_high']:.0%}" if "share_low" in d else ""
            print(f"   {d['domain']}: {d['citations']} citations ({d['share']:.0%}{rng}), in {d['runs']} of {st['runs']} runs")
    print("\nshared pages between LLMs:", summary["shared_pages_between_llms"])
    print("new names:", summary["new_names_outside_list"])
    for w in warnings + sep: print("WARNING:", w)


if __name__ == "__main__":
    main()
